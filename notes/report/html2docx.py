"""Convert the report draft's small HTML subset to .docx (headings, paragraphs,
lists, tables, inline b/i/code/sub/sup, yellow 'todo' highlights, images)."""
import sys, os
from html.parser import HTMLParser
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_COLOR_INDEX, WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

src, out = sys.argv[1], sys.argv[2]
base = os.path.dirname(os.path.abspath(src))
doc = Document()
st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(11)

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)

class P(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.para = None; self.fmt = []; self.pclass = []; self.list_stack = []
        self.table = None; self.row = None; self.cell = None; self.is_th = False; self.in_style = False
        self.rows = []
    # --- helpers
    def cur_container(self):
        return self.cell if self.cell is not None else doc
    def new_para(self, style=None):
        c = self.cur_container()
        if self.cell is not None:
            p = self.cell.paragraphs[0] if (len(self.cell.paragraphs) == 1 and not self.cell.paragraphs[0].text and not self.cell_used) else self.cell.add_paragraph()
            self.cell_used = True
        else:
            p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
        self.para = p
        return p
    def handle_starttag(self, tag, a):
        a = dict(a); cls = a.get("class", "")
        if tag == "style": self.in_style = True; return
        if tag in ("h1", "h2", "h3"):
            self.para = doc.add_heading("", level={"h1": 0, "h2": 1, "h3": 2}[tag]); self.fmt.append(("h",))
        elif tag == "p":
            self.new_para(); self.pclass.append(cls)
            if a.get("style", "").find("center") >= 0: self.para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if cls == "cap": self.para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif tag in ("ul", "ol"):
            self.list_stack.append(tag)
        elif tag == "li":
            self.para = doc.add_paragraph(style="List Bullet" if self.list_stack[-1] == "ul" else "List Number")
            self.pclass.append(cls)
        elif tag == "table":
            self.rows = []
        elif tag == "tr":
            self.rows.append([])
        elif tag in ("td", "th"):
            self.rows[-1].append({"th": tag == "th", "runs": []}); self.cellbuf = self.rows[-1][-1]["runs"]
        elif tag in ("b", "strong"): self.fmt.append(("b",))
        elif tag in ("i", "em"): self.fmt.append(("i",))
        elif tag == "code": self.fmt.append(("code",))
        elif tag == "sub": self.fmt.append(("sub",))
        elif tag == "sup": self.fmt.append(("sup",))
        elif tag == "span": self.fmt.append(("todo",) if "todo" in cls else ("span",))
        elif tag == "br":
            self.emit("\n")
        elif tag == "pre":
            self.in_pre = True; self.prebuf = []
        elif tag == "img":
            w = float(a.get("width", "6"))
            doc.add_picture(os.path.join(base, a["src"]), width=Inches(w))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    def handle_endtag(self, tag):
        if tag == "style": self.in_style = False; return
        if tag == "pre":
            self.in_pre = False
            para = doc.add_paragraph()
            lines = "".join(self.prebuf).strip("\n").split("\n")
            for k, line in enumerate(lines):
                r = para.add_run(line); r.font.name = "Courier New"; r.font.size = Pt(8.5)
                if k < len(lines) - 1: r.add_break()
            return
        if tag in ("h1", "h2", "h3"): self.fmt.pop(); self.para = None
        elif tag in ("p", "li"): self.pclass.pop(); self.para = None
        elif tag in ("ul", "ol"): self.list_stack.pop()
        elif tag in ("td", "th"): self.cellbuf = None
        elif tag == "table": self.flush_table()
        elif tag in ("b", "strong", "i", "em", "code", "sub", "sup", "span"): self.fmt.pop()
    def flags(self):
        f = {x[0] for x in self.fmt}
        pc = self.pclass[-1] if self.pclass else ""
        if pc == "todo": f.add("todo")
        if pc == "note": f.add("note")
        if pc == "cap": f.add("cap")
        return f
    def emit(self, text):
        if getattr(self, "cellbuf", None) is not None:
            self.cellbuf.append((text, self.flags())); return
        if self.para is None:
            if not text.strip(): return
            self.new_para()
        r = self.para.add_run(text); self.style_run(r, self.flags())
    def style_run(self, r, f):
        if "b" in f: r.bold = True
        if "i" in f or "note" in f: r.italic = True
        if "note" in f: r.font.color.rgb = RGBColor(0x7A, 0x5C, 0x00)
        if "cap" in f: r.font.size = Pt(9.5); r.italic = True
        if "code" in f: r.font.name = "Courier New"
        if "sub" in f: r.font.subscript = True
        if "sup" in f: r.font.superscript = True
        if "todo" in f: r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    def handle_data(self, data):
        if self.in_style: return
        if getattr(self, "in_pre", False): self.prebuf.append(data); return
        data = " ".join(data.split()) if "\n" in data or "  " in data else data
        if data == "": return
        self.emit(data)
    def flush_table(self):
        ncols = max(len(r) for r in self.rows)
        t = doc.add_table(rows=len(self.rows), cols=ncols); t.style = "Table Grid"
        for i, r in enumerate(self.rows):
            for j, c in enumerate(r):
                cell = t.cell(i, j); p = cell.paragraphs[0]
                for text, f in c["runs"]:
                    run = p.add_run(text); self.style_run(run, f); run.font.size = Pt(9.5)
                    if c["th"]: run.bold = True
                if c["th"]: shade(cell, "EEEEEE")
        doc.add_paragraph()

p = P(); p.cell = None; p.cell_used = False
p.feed(open(src, encoding="utf-8").read())
doc.save(out)
print("saved", out)
